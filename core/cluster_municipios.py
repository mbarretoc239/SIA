"""Clusterização de município (A/B/C/D por porte, planilha mantida pelo BI)
-- cruza com CIDADE/UF do REL5201 pra mostrar o cluster do prestador no
cabeçalho do processo na Amostragem (ver views/7_Amostragem_Beta.py)."""
import unicodedata
import zipfile

import openpyxl
import streamlit as st

ABA_CLUSTER = "Cluster"
COLUNAS_NECESSARIAS = {"UF", "CLUSTER"}


def _norm(texto) -> str:
    sem_acento = unicodedata.normalize("NFKD", str(texto)).encode("ASCII", "ignore").decode("ASCII")
    return sem_acento.strip().upper()


def preparar_registros_cluster(arquivo) -> list:
    """Lê a aba "Cluster" da planilha de clusterização de município (colunas
    CIDADE, CLUSTER, UF_MUN, UF, CIDADE 5201, REGIÃO) e devolve uma lista de
    dicts prontos pra gravar em cluster_municipios.

    Usa CIDADE 5201 como nome do município quando presente -- já vem no
    mesmo formato usado pelo REL5201 (ver core/relatorio_5201.py) -- e cai
    pra CIDADE quando a planilha não tiver essa coluna. Chave (cidade, uf)
    normalizada com o mesmo _norm() do REL5201 (uppercase, sem acento), pra
    bater exatamente na hora de cruzar."""
    try:
        wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    except zipfile.BadZipFile:
        raise ValueError(
            f"O arquivo '{getattr(arquivo, 'name', '')}' não é um .xlsx válido (pode estar "
            "corrompido ou o download ter sido interrompido). Baixe de novo e tente subir novamente."
        )

    if ABA_CLUSTER not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"Aba '{ABA_CLUSTER}' não encontrada na planilha "
            f"(abas disponíveis: {', '.join(wb.sheetnames)})."
        )

    ws = wb[ABA_CLUSTER]
    linhas = ws.iter_rows(values_only=True)
    header = [_norm(c) for c in next(linhas)]
    idx = {nome: i for i, nome in enumerate(header)}
    faltantes = COLUNAS_NECESSARIAS - set(idx)
    if faltantes:
        wb.close()
        raise ValueError(
            f"Colunas não encontradas na aba '{ABA_CLUSTER}': {', '.join(sorted(faltantes))}"
        )

    i_cidade_5201 = idx.get("CIDADE 5201")
    i_cidade = idx.get("CIDADE")
    i_uf = idx["UF"]
    i_cluster = idx["CLUSTER"]
    i_regiao = idx.get("REGIAO")

    # dict em vez de lista -- a planilha pode ter linhas repetidas pro mesmo
    # município (ex.: atualização parcial colada por cima); a última linha
    # de cada (cidade, uf) vence, evitando erro de PK duplicada no insert.
    vistos = {}
    for linha in linhas:
        cidade_raw = (linha[i_cidade_5201] if i_cidade_5201 is not None else None) or \
            (linha[i_cidade] if i_cidade is not None else None)
        uf_raw = linha[i_uf]
        cluster_raw = linha[i_cluster]
        if not cidade_raw or not uf_raw or not cluster_raw:
            continue
        cidade = _norm(cidade_raw)
        uf = _norm(uf_raw)
        regiao = _norm(linha[i_regiao]) if i_regiao is not None and linha[i_regiao] else None
        vistos[(cidade, uf)] = {
            "cidade": cidade,
            "uf": uf,
            "cluster": _norm(cluster_raw),
            "regiao": regiao,
        }

    wb.close()
    return list(vistos.values())


@st.cache_data(ttl=1800)
def carregar_mapa_cluster() -> dict:
    """{(cidade, uf): {"cluster": ..., "regiao": ...}} -- cacheado 30min,
    catálogo estático que só muda quando alguém reimporta a planilha em
    Configurações."""
    from shared.database import DatabaseManager
    db = DatabaseManager()
    linhas = db.buscar_cluster_municipios()
    return {(item["cidade"], item["uf"]): {"cluster": item["cluster"], "regiao": item.get("regiao")} for item in linhas}


def cluster_do_processo(cidade: str, uf: str) -> dict | None:
    """Busca o cluster pelo par (cidade, uf) já normalizado (ver _norm em
    core/relatorio_5201.py -- mesmo formato). None se algum dos dois estiver
    vazio ou não houver match no catálogo."""
    if not cidade or not uf:
        return None
    return carregar_mapa_cluster().get((cidade, uf))
